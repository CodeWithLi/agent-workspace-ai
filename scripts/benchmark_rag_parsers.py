import argparse
import asyncio
import json
import math
import subprocess
import sys
import tempfile
import time
from datetime import datetime, timezone
from importlib.metadata import PackageNotFoundError, version
from pathlib import Path
from time import perf_counter
from typing import Any
from uuid import uuid4


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from app.rag.parser_benchmark import (  # noqa: E402
    BenchmarkElement,
    BenchmarkExpectation,
    build_backend_report,
    compare_backend_reports,
    corpus_fingerprint,
    resolve_corpus_file,
    score_document,
)


def package_version(name: str) -> str:
    try:
        return version(name)
    except PackageNotFoundError:
        return "unknown"


def ratio(value: str) -> float:
    parsed = float(value)
    if not math.isfinite(parsed) or not 0 <= parsed <= 1:
        raise argparse.ArgumentTypeError("value must be between 0 and 1")
    return parsed


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
        encoding="utf-8",
    )


def read_json_object(path: Path) -> dict[str, Any]:
    def reject_constant(value: str) -> None:
        raise ValueError(f"non-finite JSON number: {value}")

    payload = json.loads(path.read_text(encoding="utf-8"), parse_constant=reject_constant)
    if not isinstance(payload, dict):
        raise ValueError("JSON root must be an object")
    return payload


def generate_corpus(output_dir: Path) -> None:
    from docx import Document
    from PIL import Image, ImageDraw, ImageFont
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    from io import BytesIO
    from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

    output_dir.mkdir(parents=True, exist_ok=True)

    document = Document()
    document.add_heading("RECOVERY MANUAL 73001", level=1)
    document.add_paragraph("POSTGRESQL RESTORE 73002 must complete before index rebuild.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Step"
    table.cell(0, 1).text = "Owner"
    table.cell(1, 0).text = "VERIFY SNAPSHOT 73003"
    table.cell(1, 1).text = "DBA"
    fixed_time = datetime(2026, 1, 1, tzinfo=timezone.utc)
    document.core_properties.created = fixed_time
    document.core_properties.modified = fixed_time
    raw_docx = BytesIO()
    document.save(raw_docx)
    raw_docx.seek(0)
    with ZipFile(raw_docx) as source, ZipFile(
        output_dir / "recovery.docx", "w", compression=ZIP_DEFLATED
    ) as destination:
        for source_info in sorted(source.infolist(), key=lambda item: item.filename):
            target_info = ZipInfo(source_info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            target_info.compress_type = ZIP_DEFLATED
            target_info.external_attr = source_info.external_attr
            target_info.create_system = source_info.create_system
            destination.writestr(target_info, source.read(source_info.filename))

    layout_pdf = canvas.Canvas(str(output_dir / "layout.pdf"), pagesize=(612, 792), invariant=1)
    layout_pdf.setFont("Helvetica-Bold", 18)
    layout_pdf.drawString(54, 735, "QUARTERLY RECOVERY PLAN 81001")
    layout_pdf.setFont("Helvetica", 12)
    layout_pdf.drawString(54, 700, "STEP ALPHA 81002 verifies the primary backup.")
    x_positions = (54, 280, 500)
    y_positions = (650, 620, 590)
    for x in x_positions:
        layout_pdf.line(x, y_positions[-1], x, y_positions[0])
    for y in y_positions:
        layout_pdf.line(x_positions[0], y, x_positions[-1], y)
    layout_pdf.drawString(65, 630, "Step")
    layout_pdf.drawString(295, 630, "Owner")
    layout_pdf.drawString(65, 600, "TABLE CHECK 81003")
    layout_pdf.drawString(295, 600, "SRE")
    layout_pdf.showPage()
    layout_pdf.setFont("Helvetica-Bold", 16)
    layout_pdf.drawString(54, 735, "SECOND PAGE 81004")
    layout_pdf.setFont("Helvetica", 12)
    layout_pdf.drawString(54, 700, "STEP BETA 81005 validates the restore drill.")
    layout_pdf.save()

    image = Image.new("RGB", (1600, 500), "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 72)
    except OSError:
        font = ImageFont.load_default(size=72)
    draw.text((80, 100), "OCR RECOVERY CODE 48271", fill="black", font=font)
    png = BytesIO()
    image.save(png, format="PNG")
    png.seek(0)
    scanned_pdf = canvas.Canvas(str(output_dir / "scanned.pdf"), pagesize=(612, 792), invariant=1)
    scanned_pdf.drawImage(ImageReader(png), 36, 300, width=540, height=168, preserveAspectRatio=True)
    scanned_pdf.showPage()
    scanned_pdf.save()

    manifest = {
        "schema_version": 1,
        "samples": [
            {
                "filename": "recovery.docx",
                "markers": ["RECOVERY MANUAL 73001", "POSTGRESQL RESTORE 73002", "VERIFY SNAPSHOT 73003"],
                "ordered_markers": [
                    "RECOVERY MANUAL 73001",
                    "POSTGRESQL RESTORE 73002",
                    "VERIFY SNAPSHOT 73003",
                ],
                "required_kinds": ["title", "table"],
            },
            {
                "filename": "layout.pdf",
                "markers": [
                    "QUARTERLY RECOVERY PLAN 81001",
                    "STEP ALPHA 81002",
                    "TABLE CHECK 81003",
                    "SECOND PAGE 81004",
                    "STEP BETA 81005",
                ],
                "ordered_markers": [
                    "QUARTERLY RECOVERY PLAN 81001",
                    "STEP ALPHA 81002",
                    "TABLE CHECK 81003",
                    "SECOND PAGE 81004",
                    "STEP BETA 81005",
                ],
                "required_kinds": ["title", "table"],
                "require_page_metadata": True,
                "require_bbox_metadata": True,
            },
            {
                "filename": "scanned.pdf",
                "markers": ["OCR RECOVERY CODE 48271"],
                "ordered_markers": ["OCR RECOVERY CODE 48271"],
                "required_kinds": ["paragraph"],
                "require_page_metadata": True,
                "require_bbox_metadata": True,
            },
        ],
    }
    write_json(output_dir / "manifest.json", manifest)


def load_expectations(corpus_dir: Path) -> list[BenchmarkExpectation]:
    payload = json.loads((corpus_dir / "manifest.json").read_text(encoding="utf-8"))
    if not isinstance(payload, dict) or payload.get("schema_version") != 1:
        raise ValueError("benchmark manifest has an unsupported schema")
    rows = payload.get("samples")
    if not isinstance(rows, list) or not rows:
        raise ValueError("benchmark manifest must contain samples")
    expectations: list[BenchmarkExpectation] = []
    for index, row in enumerate(rows):
        if not isinstance(row, dict):
            raise ValueError(f"benchmark sample {index} must be an object")
        filename = row.get("filename")
        markers = row.get("markers")
        ordered_markers = row.get("ordered_markers")
        required_kinds = row.get("required_kinds")
        if not isinstance(filename, str) or not filename:
            raise ValueError(f"benchmark sample {index} has no filename")
        for field, value in (
            ("markers", markers),
            ("ordered_markers", ordered_markers),
            ("required_kinds", required_kinds),
        ):
            if not isinstance(value, list) or not value or not all(
                isinstance(item, str) and item for item in value
            ):
                raise ValueError(f"benchmark sample {filename} has invalid {field}")
        resolve_corpus_file(corpus_dir, filename)
        expectations.append(
            BenchmarkExpectation(
                filename=filename,
                markers=tuple(markers),
                ordered_markers=tuple(ordered_markers),
                required_kinds=tuple(required_kinds),
                require_page_metadata=bool(row.get("require_page_metadata", False)),
                require_bbox_metadata=bool(row.get("require_bbox_metadata", False)),
            )
        )
    return expectations


def as_bbox(value: Any) -> tuple[float, float, float, float] | None:
    if isinstance(value, (list, tuple)) and len(value) == 4 and all(
        isinstance(item, (int, float)) for item in value
    ):
        return tuple(float(item) for item in value)
    return None


def parsed_document_elements(path: Path) -> tuple[list[BenchmarkElement], str]:
    from app.rag.parsers import parse_document_elements

    parsed = parse_document_elements(path.name, path.read_bytes())
    elements = [
        BenchmarkElement(
            kind=element.kind,
            content=element.content,
            page=element.page_start,
            bbox=as_bbox(element.metadata.get("bbox")),
        )
        for element in parsed.elements
    ]
    return elements, f"{parsed.parser_name}-{parsed.parser_version}"


async def docling_elements(path: Path) -> tuple[list[BenchmarkElement], str]:
    from app.rag.advanced_parser import DoclingServeParser

    parsed = await DoclingServeParser().parse(path.name, path.read_bytes())
    return (
        [
            BenchmarkElement(
                kind=element.kind,
                content=element.content,
                page=element.page_start,
                bbox=as_bbox(element.metadata.get("bbox")),
            )
            for element in parsed.elements
        ],
        parsed.parser_version,
    )


def unstructured_elements(path: Path) -> tuple[list[BenchmarkElement], str]:
    from unstructured.partition.auto import partition

    kwargs: dict[str, Any] = {}
    if path.suffix.lower() == ".pdf":
        kwargs = {"strategy": "hi_res", "infer_table_structure": True, "languages": ["eng"]}
    raw_elements = partition(filename=str(path), **kwargs)
    kind_map = {
        "Title": "title",
        "Table": "table",
        "CodeSnippet": "code",
        "ListItem": "list_item",
    }
    elements: list[BenchmarkElement] = []
    for element in raw_elements:
        metadata = element.metadata.to_dict() if getattr(element, "metadata", None) else {}
        category = str(getattr(element, "category", element.__class__.__name__))
        content = str(element)
        if category == "Table" and metadata.get("text_as_html"):
            content += "\n" + str(metadata["text_as_html"])
        coordinates = metadata.get("coordinates") or {}
        points = coordinates.get("points") if isinstance(coordinates, dict) else None
        bbox = None
        if isinstance(points, (list, tuple)) and points:
            xs = [point[0] for point in points if isinstance(point, (list, tuple)) and len(point) == 2]
            ys = [point[1] for point in points if isinstance(point, (list, tuple)) and len(point) == 2]
            if xs and ys:
                bbox = (float(min(xs)), float(min(ys)), float(max(xs)), float(max(ys)))
        page_number = metadata.get("page_number")
        elements.append(
            BenchmarkElement(
                kind=kind_map.get(category, "paragraph"),
                content=content,
                page=int(page_number)
                if isinstance(page_number, (int, float)) and not isinstance(page_number, bool)
                else None,
                bbox=bbox,
            )
        )
    return elements, package_version("unstructured")


def nested_strings(value: Any) -> list[str]:
    if isinstance(value, str):
        return [value]
    if isinstance(value, list):
        return [text for item in value for text in nested_strings(item)]
    if isinstance(value, dict):
        return [
            text
            for key, item in value.items()
            if key not in {"img_path", "image_path", "bbox"}
            for text in nested_strings(item)
        ]
    return []


def mineru_content_list_elements(path: Path) -> list[BenchmarkElement]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    items: list[tuple[int, dict[str, Any]]] = []
    if isinstance(payload, list) and payload and all(isinstance(page, list) for page in payload):
        items = [(page_index, item) for page_index, page in enumerate(payload) for item in page if isinstance(item, dict)]
    elif isinstance(payload, list):
        items = [
            (int(item.get("page_idx", 0)), item)
            for item in payload
            if isinstance(item, dict)
        ]
    kind_map = {
        "title": "title",
        "table": "table",
        "code": "code",
        "algorithm": "code",
        "list": "list_item",
        "text": "paragraph",
        "paragraph": "paragraph",
    }
    elements: list[BenchmarkElement] = []
    for page_index, item in items:
        item_type = str(item.get("type", "paragraph"))
        if item_type == "text" and int(item.get("text_level", 0) or 0) > 0:
            kind = "title"
        else:
            kind = kind_map.get(item_type, "paragraph")
        content = "\n".join(nested_strings(item.get("content", item))).strip()
        elements.append(
            BenchmarkElement(
                kind=kind,
                content=content,
                page=page_index + 1,
                bbox=as_bbox(item.get("bbox")),
            )
        )
    return elements


def peak_rss_mb(children: bool = False) -> float | None:
    try:
        import resource

        usage = resource.getrusage(resource.RUSAGE_CHILDREN if children else resource.RUSAGE_SELF)
        divisor = 1024 * 1024 if sys.platform == "darwin" else 1024
        return float(usage.ru_maxrss) / divisor
    except (ImportError, AttributeError):
        return None


def sampled_peak_memory_mb(path: Path | None) -> float | None:
    if path is None or not path.exists():
        return None
    samples: list[float] = []
    for line in path.read_text(encoding="utf-8").splitlines():
        try:
            samples.append(float(line.strip()) / (1024 * 1024))
        except ValueError:
            continue
    return max(samples, default=None)


async def run_regular_backend(
    backend: str,
    corpus_dir: Path,
    expectations: list[BenchmarkExpectation],
) -> tuple[list[dict[str, Any]], str]:
    rows: list[dict[str, Any]] = []
    parser_versions: set[str] = set()
    for expectation in expectations:
        started = perf_counter()
        try:
            if backend == "native":
                elements, parser_version = parsed_document_elements(
                    resolve_corpus_file(corpus_dir, expectation.filename)
                )
            elif backend == "docling":
                elements, parser_version = await docling_elements(
                    resolve_corpus_file(corpus_dir, expectation.filename)
                )
            elif backend == "unstructured":
                elements, parser_version = unstructured_elements(
                    resolve_corpus_file(corpus_dir, expectation.filename)
                )
            else:
                raise ValueError(f"Unsupported regular backend: {backend}")
            error = ""
            parser_versions.add(parser_version)
        except Exception as exc:
            elements = []
            error = f"{type(exc).__name__}: {exc}"
        rows.append(
            score_document(
                expectation,
                elements,
                latency_ms=(perf_counter() - started) * 1000,
                error=error,
            )
        )
    return rows, ",".join(sorted(parser_versions)) or "unknown"


def run_mineru_backend(
    corpus_dir: Path,
    expectations: list[BenchmarkExpectation],
    output_dir: Path,
    timeout_seconds: int,
    api_url: str | None,
) -> tuple[list[dict[str, Any]], str, float | None]:
    command = [
        "mineru",
        "-p",
        str(corpus_dir),
        "-o",
        str(output_dir),
        "-b",
        "pipeline",
        "-m",
        "auto",
        "-l",
        "en",
        "-t",
        "true",
        "-f",
        "true",
    ]
    if api_url:
        command.extend(["--api-url", api_url])
    started = perf_counter()
    peak_memory_mb: float | None = None
    try:
        import psutil

        with tempfile.TemporaryFile(mode="w+", encoding="utf-8") as log_file:
            process = subprocess.Popen(command, stdout=log_file, stderr=log_file, text=True)
            deadline = time.monotonic() + timeout_seconds
            while process.poll() is None:
                try:
                    root_process = psutil.Process(process.pid)
                    processes = [root_process, *root_process.children(recursive=True)]
                    rss = sum(item.memory_info().rss for item in processes if item.is_running())
                    peak_memory_mb = max(peak_memory_mb or 0.0, rss / (1024 * 1024))
                except (psutil.Error, OSError):
                    pass
                if time.monotonic() >= deadline:
                    try:
                        for child in psutil.Process(process.pid).children(recursive=True):
                            child.kill()
                    except psutil.Error:
                        pass
                    process.kill()
                    process.wait()
                    raise subprocess.TimeoutExpired(command, timeout_seconds)
                time.sleep(0.25)
            log_file.seek(0)
            command_output = log_file.read()
            command_error = "" if process.returncode == 0 else (command_output or "MinerU command failed")[-1000:]
    except (ImportError, OSError, subprocess.TimeoutExpired) as exc:
        command_error = f"{type(exc).__name__}: {exc}"
    elapsed_ms = (perf_counter() - started) * 1000
    rows: list[dict[str, Any]] = []
    per_document_latency = elapsed_ms / max(len(expectations), 1)
    for expectation in expectations:
        expected_name = f"{Path(expectation.filename).stem}_content_list.json"
        matches = sorted(
            path for path in output_dir.rglob("*_content_list.json") if path.name == expected_name
        )
        error = command_error
        elements: list[BenchmarkElement] = []
        if not error:
            if len(matches) != 1:
                error = f"expected one content list, found {len(matches)}"
            else:
                try:
                    elements = mineru_content_list_elements(matches[0])
                except Exception as exc:
                    error = f"{type(exc).__name__}: {exc}"
        rows.append(
            score_document(
                expectation,
                elements,
                latency_ms=per_document_latency,
                error=error,
            )
        )
    return rows, package_version("mineru"), peak_memory_mb


async def run_backend(args: argparse.Namespace) -> int:
    expectations = load_expectations(args.corpus_dir)
    fingerprint = corpus_fingerprint(args.corpus_dir, expectations)
    if args.backend == "mineru":
        mineru_output = args.work_dir / f"mineru-{uuid4().hex}"
        mineru_output.mkdir(parents=True)
        rows, backend_version, measured_peak = run_mineru_backend(
            args.corpus_dir,
            expectations,
            mineru_output,
            args.timeout_seconds,
            args.mineru_api_url,
        )
        process_peak = measured_peak if measured_peak is not None else peak_rss_mb(children=True)
    else:
        rows, backend_version = await run_regular_backend(args.backend, args.corpus_dir, expectations)
        process_peak = peak_rss_mb()
    sampled_peak = sampled_peak_memory_mb(args.resource_sample_file)
    report = build_backend_report(
        backend=args.backend,
        version=backend_version,
        corpus_hash=fingerprint,
        documents=rows,
        peak_memory_mb=sampled_peak if sampled_peak is not None else process_peak,
        artifact_size_mb=args.artifact_size_mb,
        memory_scope=args.memory_scope,
        artifact_scope=args.artifact_scope,
    )
    write_json(args.output, report)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


def compare(args: argparse.Namespace) -> int:
    report_paths = list(args.reports)
    if args.report_dir and args.report_dir.exists():
        report_paths.extend(sorted(args.report_dir.glob("*.json")))
    unique_paths = sorted(set(report_paths))
    reports: list[dict[str, Any]] = []
    load_failures: list[str] = []
    for path in unique_paths:
        try:
            reports.append(read_json_object(path))
        except (OSError, json.JSONDecodeError, ValueError) as exc:
            load_failures.append(f"{path}: {exc}")
    payload = compare_backend_reports(
        reports,
        required_backends=tuple(args.required_backends),
        min_success_rate=args.min_success_rate,
        min_marker_recall=args.min_marker_recall,
        min_structure_recall=args.min_structure_recall,
        min_reading_order_accuracy=args.min_reading_order_accuracy,
        min_page_metadata_coverage=args.min_page_metadata_coverage,
        min_bbox_metadata_coverage=args.min_bbox_metadata_coverage,
    )
    if load_failures:
        payload["failures"] = [*load_failures, *payload["failures"]]
        payload["passed"] = False
    write_json(args.output, payload)
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    return 1 if args.enforce_gates and not payload["passed"] else 0


def annotate(args: argparse.Namespace) -> int:
    payload = read_json_object(args.report)
    metrics = payload.get("metrics")
    if not isinstance(metrics, dict):
        raise ValueError("Parser report does not contain metrics")
    if args.peak_memory_mb is not None:
        if not math.isfinite(args.peak_memory_mb) or args.peak_memory_mb <= 0:
            raise ValueError("peak memory must be a positive finite number")
        metrics["peak_memory_mb"] = round(args.peak_memory_mb, 2)
    if args.artifact_size_mb is not None:
        if not math.isfinite(args.artifact_size_mb) or args.artifact_size_mb <= 0:
            raise ValueError("artifact size must be a positive finite number")
        metrics["artifact_size_mb"] = round(args.artifact_size_mb, 2)
    if args.memory_scope:
        payload["memory_scope"] = args.memory_scope
    if args.artifact_scope:
        payload["artifact_scope"] = args.artifact_scope
    write_json(args.report, payload)
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description="Run a same-corpus RAG parser benchmark")
    subparsers = parser.add_subparsers(dest="command", required=True)

    generate_parser = subparsers.add_parser("generate")
    generate_parser.add_argument("--output-dir", type=Path, required=True)

    run_parser = subparsers.add_parser("run")
    run_parser.add_argument("--backend", choices=("native", "docling", "unstructured", "mineru"), required=True)
    run_parser.add_argument("--corpus-dir", type=Path, required=True)
    run_parser.add_argument("--output", type=Path, required=True)
    run_parser.add_argument("--work-dir", type=Path, default=Path(".run/parser-work"))
    run_parser.add_argument("--timeout-seconds", type=int, default=7200)
    run_parser.add_argument("--mineru-api-url")
    run_parser.add_argument("--resource-sample-file", type=Path)
    run_parser.add_argument("--artifact-size-mb", type=float)
    run_parser.add_argument("--memory-scope", default="process")
    run_parser.add_argument("--artifact-scope", default="unknown")

    compare_parser = subparsers.add_parser("compare")
    compare_parser.add_argument("--reports", type=Path, nargs="*", default=[])
    compare_parser.add_argument("--report-dir", type=Path)
    compare_parser.add_argument("--output", type=Path, required=True)
    compare_parser.add_argument("--enforce-gates", action="store_true")
    compare_parser.add_argument(
        "--required-backends",
        nargs="+",
        choices=("native", "docling", "unstructured", "mineru"),
        default=("native", "docling", "unstructured", "mineru"),
    )
    compare_parser.add_argument("--min-success-rate", type=ratio, default=1.0)
    compare_parser.add_argument("--min-marker-recall", type=ratio, default=0.9)
    compare_parser.add_argument("--min-structure-recall", type=ratio, default=0.5)
    compare_parser.add_argument("--min-reading-order-accuracy", type=ratio, default=1.0)
    compare_parser.add_argument("--min-page-metadata-coverage", type=ratio, default=0.9)
    compare_parser.add_argument("--min-bbox-metadata-coverage", type=ratio, default=0.9)

    annotate_parser = subparsers.add_parser("annotate")
    annotate_parser.add_argument("--report", type=Path, required=True)
    annotate_parser.add_argument("--peak-memory-mb", type=float)
    annotate_parser.add_argument("--artifact-size-mb", type=float)
    annotate_parser.add_argument("--memory-scope")
    annotate_parser.add_argument("--artifact-scope")

    args = parser.parse_args()
    if args.command == "generate":
        generate_corpus(args.output_dir)
        return 0
    if args.command == "run":
        args.work_dir.mkdir(parents=True, exist_ok=True)
        return asyncio.run(run_backend(args))
    if args.command == "compare":
        return compare(args)
    return annotate(args)


if __name__ == "__main__":
    raise SystemExit(main())
