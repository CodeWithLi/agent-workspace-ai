import os
from io import BytesIO

import pytest
from docx import Document

from app.rag.advanced_parser import DoclingServeParser, settings as docling_settings


pytestmark = pytest.mark.skipif(
    os.getenv("RAG_REAL_DOCLING") != "1",
    reason="requires a real Docling Serve instance",
)


def scanned_pdf() -> bytes:
    from PIL import Image, ImageDraw, ImageFont
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    image = Image.new("RGB", (1600, 500), "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 72)
    except OSError:
        font = ImageFont.load_default()
    draw.text((80, 100), "OCR RECOVERY CODE 48271", fill="black", font=font)
    png = BytesIO()
    image.save(png, format="PNG")
    png.seek(0)

    pdf = BytesIO()
    document = canvas.Canvas(pdf, pagesize=(612, 792))
    document.drawImage(ImageReader(png), 36, 300, width=540, height=168, preserveAspectRatio=True)
    document.showPage()
    document.save()
    return pdf.getvalue()


@pytest.mark.asyncio
async def test_real_docling_preserves_docx_table_and_ocr_text() -> None:
    assert docling_settings.docling_serve_url
    assert docling_settings.docling_do_ocr is True
    parser = DoclingServeParser()
    await parser.healthcheck()

    document = Document()
    document.add_heading("Recovery Manual", level=1)
    document.add_paragraph("Restore PostgreSQL before rebuilding the Milvus collection.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Step"
    table.cell(0, 1).text = "Owner"
    table.cell(1, 0).text = "Restore"
    table.cell(1, 1).text = "DBA"
    stream = BytesIO()
    document.save(stream)

    parsed_docx = await parser.parse("recovery.docx", stream.getvalue())
    parsed_pdf = await parser.parse("scanned-recovery.pdf", scanned_pdf())

    assert parsed_docx.parser_name == "docling-serve"
    assert "Recovery Manual" in parsed_docx.text
    assert "PostgreSQL" in parsed_docx.text
    assert any(element.kind == "table" and "Owner" in element.content for element in parsed_docx.elements)
    assert parsed_pdf.parser_name == "docling-serve"
    assert "48271" in parsed_pdf.text
