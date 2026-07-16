import json
import math
import re
from collections import Counter
from io import BytesIO
from pathlib import Path
from uuid import uuid4

from fastapi import UploadFile
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.core.llm import get_llm
from app.core.storage import object_storage
from app.models import Embedding, KnowledgeDocument, User
from app.rag.vector_store import vector_store


settings = get_settings()
ALLOWED_SUFFIXES = {".pdf", ".docx", ".md", ".markdown", ".txt"}


class DocumentValidationError(ValueError):
    pass


async def read_upload(file: UploadFile) -> bytes:
    limit = settings.max_upload_mb * 1024 * 1024
    data = await file.read(limit + 1)
    if len(data) > limit:
        raise DocumentValidationError(f"文件大小不能超过 {settings.max_upload_mb} MB")
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in ALLOWED_SUFFIXES:
        raise DocumentValidationError("仅支持 PDF、DOCX、Markdown 和 TXT 文件")
    return data


def parse_document(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(BytesIO(data))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        if suffix == ".docx":
            from docx import Document

            document = Document(BytesIO(data))
            return "\n".join(paragraph.text for paragraph in document.paragraphs)
        return data.decode("utf-8")
    except Exception as exc:
        raise DocumentValidationError("文件解析失败，请确认文件未损坏且编码正确") from exc


def split_text(text: str, chunk_size: int = 900, overlap: int = 140) -> list[str]:
    clean = text.strip()
    if not clean:
        return []
    chunks: list[str] = []
    start = 0
    while start < len(clean):
        end = min(start + chunk_size, len(clean))
        if end < len(clean):
            boundary = max(clean.rfind("\n", start, end), clean.rfind("。", start, end))
            if boundary > start + chunk_size // 2:
                end = boundary + 1
        chunks.append(clean[start:end].strip())
        if end == len(clean):
            break
        start = max(0, end - overlap)
    return [chunk for chunk in chunks if chunk]


async def index_document(
    db: Session,
    user: User,
    file: UploadFile,
    project_id: int | None = None,
) -> KnowledgeDocument:
    filename = Path(file.filename or "document").name
    data = await read_upload(file)
    text = parse_document(filename, data)
    chunks = split_text(text)
    object_name = f"users/{user.id}/{uuid4().hex}_{filename}"
    storage_path = object_storage.save(object_name, data, file.content_type or "application/octet-stream")
    document = KnowledgeDocument(
        project_id=project_id,
        owner_id=user.id,
        filename=filename,
        content_type=file.content_type or "application/octet-stream",
        storage_path=storage_path,
        status="INDEXING" if chunks else "EMPTY",
        summary=text[:500],
    )
    db.add(document)
    db.flush()

    try:
        vectors = await get_llm().embed_many(chunks)
        for index, (content, vector) in enumerate(zip(chunks, vectors)):
            vector_id = f"chunk-{document.id}-{index}"
            vector_store.upsert(vector_id, document.id, project_id, content, vector)
            db.add(
                Embedding(
                    document_id=document.id,
                    project_id=project_id,
                    content=content,
                    chunk_index=index,
                    vector_id=vector_id,
                    vector_json=json.dumps(vector),
                    score_hint=sum(abs(value) for value in vector) / max(len(vector), 1),
                )
            )
        if chunks:
            document.status = "INDEXED"
        db.commit()
        db.refresh(document)
        return document
    except Exception:
        db.rollback()
        object_storage.delete(storage_path)
        raise


def tokenize(text: str) -> list[str]:
    lower = text.lower()
    latin = re.findall(r"[a-z0-9_]+", lower)
    cjk = "".join(re.findall(r"[\u4e00-\u9fff]", lower))
    return latin + list(cjk) + [cjk[index : index + 2] for index in range(len(cjk) - 1)]


def sparse_scores(query: str, contents: list[str]) -> list[float]:
    query_terms = set(tokenize(query))
    tokenized = [tokenize(content) for content in contents]
    if not query_terms or not tokenized:
        return [0.0] * len(contents)
    document_frequency = Counter(
        term for terms in tokenized for term in query_terms if term in set(terms)
    )
    average_length = sum(len(terms) for terms in tokenized) / max(len(tokenized), 1)
    raw_scores: list[float] = []
    for terms in tokenized:
        frequencies = Counter(terms)
        score = 0.0
        for term in query_terms:
            frequency = frequencies.get(term, 0)
            if not frequency:
                continue
            idf = math.log(1 + (len(tokenized) - document_frequency[term] + 0.5) / (document_frequency[term] + 0.5))
            denominator = frequency + 1.5 * (0.25 + 0.75 * len(terms) / max(average_length, 1))
            score += idf * frequency * 2.5 / denominator
        raw_scores.append(score)
    maximum = max(raw_scores, default=0.0)
    return [score / maximum if maximum else 0.0 for score in raw_scores]


def sparse_score(query: str, content: str) -> float:
    return sparse_scores(query, [content])[0]


def cosine_similarity(left: list[float], right: list[float]) -> float:
    dot = sum(a * b for a, b in zip(left, right))
    left_norm = math.sqrt(sum(a * a for a in left))
    right_norm = math.sqrt(sum(b * b for b in right))
    if not left_norm or not right_norm:
        return 0.0
    return max(0.0, min(1.0, dot / (left_norm * right_norm)))


def _rank_map(rows: list[dict], key: str) -> dict[str, int]:
    ranked = sorted(rows, key=lambda item: item[key], reverse=True)
    return {row["vector_id"]: index + 1 for index, row in enumerate(ranked)}


def rerank(query: str, rows: list[dict]) -> list[dict]:
    query_tokens = set(tokenize(query))
    for row in rows:
        content_tokens = set(tokenize(row["content"]))
        coverage = len(query_tokens & content_tokens) / max(len(query_tokens), 1)
        row["score"] = min(1.0, row["score"] + coverage * 0.08)
    return sorted(rows, key=lambda item: item["score"], reverse=True)


async def search_knowledge(
    db: Session,
    user: User,
    query: str,
    project_id: int | None = None,
    top_k: int = 5,
    min_score: float = 0.2,
) -> tuple[str, list[dict]]:
    stmt = (
        select(Embedding, KnowledgeDocument)
        .join(KnowledgeDocument, KnowledgeDocument.id == Embedding.document_id)
        .where(KnowledgeDocument.status == "INDEXED")
    )
    if project_id is not None:
        stmt = stmt.where(Embedding.project_id == project_id)
    else:
        stmt = stmt.where(KnowledgeDocument.owner_id == user.id)
    records = list(db.execute(stmt))
    if not records:
        return "未检索到可用的知识库内容，请先上传并完成索引。", []

    query_vector = await get_llm().embed(query)
    stored_vectors: list[list[float]] = []
    missing_indexes: list[int] = []
    for index, (chunk, _) in enumerate(records):
        try:
            vector = json.loads(chunk.vector_json or "[]")
        except json.JSONDecodeError:
            vector = []
        stored_vectors.append(vector)
        if not vector:
            missing_indexes.append(index)
    if missing_indexes:
        generated = await get_llm().embed_many([records[index][0].content for index in missing_indexes])
        for index, vector in zip(missing_indexes, generated):
            stored_vectors[index] = vector

    sparse = sparse_scores(query, [chunk.content for chunk, _ in records])
    ann_rows = vector_store.search(query_vector, project_id, limit=max(top_k * 8, 40))
    ann_scores = {row["vector_id"]: row["score"] for row in ann_rows}
    rows: list[dict] = []
    for index, ((chunk, document), vector) in enumerate(zip(records, stored_vectors)):
        dense = ann_scores.get(chunk.vector_id or "", cosine_similarity(query_vector, vector))
        rows.append(
            {
                "vector_id": chunk.vector_id or str(chunk.id),
                "document_id": document.id,
                "filename": document.filename,
                "content": chunk.content,
                "citation": f"{document.filename}#chunk-{chunk.chunk_index}",
                "sparse_score": round(sparse[index], 4),
                "dense_score": round(dense, 4),
            }
        )

    sparse_ranks = _rank_map(rows, "sparse_score")
    dense_ranks = _rank_map(rows, "dense_score")
    for row in rows:
        sparse_rank_score = 1 / (1 + (sparse_ranks[row["vector_id"]] - 1) / 10)
        dense_rank_score = 1 / (1 + (dense_ranks[row["vector_id"]] - 1) / 10)
        fusion = 0.45 * sparse_rank_score + 0.55 * dense_rank_score
        row["score"] = 0.25 * row["sparse_score"] + 0.4 * row["dense_score"] + 0.35 * fusion

    results = [row for row in rerank(query, rows) if row["score"] >= min_score][:top_k]
    for row in results:
        row["score"] = round(row["score"], 4)
        row.pop("vector_id", None)
    if not results:
        return "未找到达到相关度阈值的知识内容，请调整关键词或降低相似度阈值。", []

    context = "\n\n".join(f"[{item['citation']}]\n{item['content']}" for item in results)
    answer = await get_llm().chat(
        "你是AI知识库管理员。只根据给定上下文回答；每个关键结论都要使用 [文件名#chunk-N] 引用。上下文不足时明确说明。",
        f"问题：{query}\n\n上下文：{context}",
    )
    return answer, results
